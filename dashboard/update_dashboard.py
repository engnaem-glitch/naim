#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
سكربت تحديث لوحة سير العمل — المساعدة النقدية لإصلاح المأوى الطارئ (25 مستفيدًا)

يقرأ ملف Google Sheets الموحّد (بعد تصديره xlsx) ويستبدل كتل البيانات فقط داخل HTML اللوحة،
ثم يتحقق من كل شيء قبل كتابة الناتج. لا يلمس CSS ولا HTML ولا منطق العرض ولا MILESTONES/COORDS/PLAN.

الاستخدام:
  python3 update_dashboard.py --xlsx src.xlsx --html current.html --out new.html [--file-id ID]

رمز الخروج: 0 = نجح التحقق والملف جاهز للنشر · 2 = فشل التحقق (لا تنشر) · 1 = خطأ غير متوقع.
يطبع في النهاية سطر JSON واحدًا يبدأ بـ RESULT: يلخّص التحديث (للمهمة المجدولة).

قواعد ثابتة (معتمدة من المستخدم):
- لا رقم هوية ولا رقم جوال ولا أي قيمة مالية خاصة بمستفيد بعينه في اللوحة.
- توزيع المهندسات ثابت حسب الكود (ENG_MAP) ولا يُقرأ من عمود "المهندس المشرف".
- FINANCE.totalGrant = 274998 ثابت. المصروف فعليًا من ورقة "ملخص المستفيدين المرتبط" (أعمدة "حسب الصرف").
- الأعمدة تُحدَّد بأسماء عناوينها في الصف 3 وليس بمواقعها، لأن المصدر يُعاد ترتيبه أحيانًا.
"""
import argparse, json, re, subprocess, sys, tempfile, os
from datetime import datetime, date, timedelta, timezone

try:
    from openpyxl import load_workbook
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "-q", "openpyxl"])
    from openpyxl import load_workbook

TOTAL_GRANT = 274998
EXPECTED_N = 25
GAZA_OFFSET = timedelta(hours=3)  # توقيت فلسطين الصيفي (UTC+3)؛ يكفي لتحديد تاريخ اليوم وساعة المزامنة

ENG_MAP = {}
for code in "MM4 NS1 NS2 RM11 RM5 MM5 MM6 MM8 MM9 NS9 MF7 MS8".split():
    ENG_MAP[code] = "م. رانيا مهنا"
for code in "MR6 MM2 MR3 MS5 MS1 MR2 MF8 MF6 MF5 SE5 SE6 SE7 SE11".split():
    ENG_MAP[code] = "م. سماح ابراهيم"

CODE_RE = re.compile(r"^[A-Z]{2}\d{1,2}$")
FIELDS = ["code", "name", "eng", "items", "completion", "p1", "p2", "schedGood", "eligGood", "housing",
          "schedText", "eligText", "stage", "adherence", "familySize", "hasExtended", "disabled",
          "chronic", "warInjured", "elderly"]


class ValidationError(Exception):
    pass


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def num(v, default=0.0):
    if v is None or v == "":
        return default
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return default


def as_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)) and 40000 < v < 60000:  # رقم تسلسلي لتاريخ Excel
        return (date(1899, 12, 30) + timedelta(days=int(v))).isoformat()
    m = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", str(v))
    if m:
        return date(int(m[1]), int(m[2]), int(m[3])).isoformat()
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", str(v))
    if m:
        return date(int(m[3]), int(m[2]), int(m[1])).isoformat()
    return None


class Sheet:
    """ورقة بصف عناوين (افتراضيًا الصف 3) وصفوف مستفيدين مفهرسة بكود المستفيد."""

    def __init__(self, wb, name, header_row=3):
        if name not in wb.sheetnames:
            raise ValidationError(f"الورقة غير موجودة في الملف المصدر: {name}")
        self.ws = wb[name]
        self.name = name
        self.headers = {}
        for c in range(1, self.ws.max_column + 1):
            h = norm(self.ws.cell(header_row, c).value)
            if h:
                self.headers[c] = h
        self.code_col = self.col("كود المستفيد")
        self.rows = {}
        for r in range(header_row + 1, self.ws.max_row + 1):
            code = norm(self.ws.cell(r, self.code_col).value)
            if CODE_RE.match(code):
                if code in self.rows:
                    raise ValidationError(f"كود مكرر {code} في ورقة {name}")
                self.rows[code] = r

    def col(self, *needles, required=True):
        """أول عمود يحوي عنوانه كل الكلمات المطلوبة."""
        for c, h in self.headers.items():
            if all(n in h for n in needles):
                return c
        if required:
            raise ValidationError(f"لم يُعثر على عمود يحوي {needles} في ورقة {self.name}")
        return None

    def get(self, code, c):
        return self.ws.cell(self.rows[code], c).value


def find_label(ws, text, rows=None):
    """موضع أول خلية تبدأ بالنص المطلوب."""
    rng = rows or range(1, ws.max_row + 1)
    for r in rng:
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and norm(v).startswith(text):
                return r, c
    return None


def read_source(xlsx):
    wb = load_workbook(xlsx, data_only=True)
    summ = Sheet(wb, "ملخص المستفيدين المرتبط")
    flow = Sheet(wb, "لوحة سير المشروع")
    fam = Sheet(wb, "قاعدة بيانات الأسر")

    c_name = summ.col("اسم المستفيد")
    c_items = summ.col("عدد بنود العقد")
    c_housing = summ.col("حاله السكن")
    c_comp = summ.col("نسبة الانجاز الكلي")
    c_p1 = summ.col("نسبة انجاز الدفعه الاولى")
    c_p2 = summ.col("نسبة انجاز الدفعه الثانية")
    c_elig = summ.col("مؤهل للدفعه")
    c_sched = summ.col("الحالة الزمنية")
    c_rcv = [summ.col("استلام الدفعه الاولى"), summ.col("استلام الدفعه الثانية"), summ.col("استلام الدفعه الثالثه")]
    c_paid = [summ.col("قيمة الدفعة الأولى حسب الصرف"), summ.col("قيمة الدفعة الثانية حسب الصرف"),
              summ.col("قيمة الدفعة الثالثة حسب الصرف")]
    c_paid_total = summ.col("إجمالي المصروف")

    f_stage = flow.col("المرحلة الحالية")
    f_adh = flow.col("حالة المتابعة الزمنية")
    f_contract = flow.col("القيمة الإجمالية للعقد")

    a_core = fam.col("عدد أفراد الأسرة الأساسية")
    a_dis = fam.col("ذوي الإعاقة")
    a_chr = fam.col("الأمراض المزمنة")
    a_war = fam.col("المصابين في الحرب")
    a_nurse = fam.col("المرضعات")
    a_hasext = fam.col("توجد أسرة ممتدة")
    a_subfam = fam.col("عدد الأسر الممتدة الفرعية")
    a_extm = fam.col("إجمالي أفراد الأسر الممتدة")
    a_eld = fam.col("إجمالي كبار السن")
    a_total = fam.col("إجمالي أفراد الوحدة السكنية")

    codes = list(summ.rows.keys())
    unknown = [c for c in codes if c not in ENG_MAP]
    if unknown:
        raise ValidationError("أكواد مستفيدين جديدة غير موجودة في جدول توزيع المهندسات: " + ", ".join(unknown)
                              + " — يلزم قرار المستخدم قبل النشر")
    for c in codes:
        if c not in flow.rows:
            raise ValidationError(f"الكود {c} غير موجود في ورقة لوحة سير المشروع")
        if c not in fam.rows:
            raise ValidationError(f"الكود {c} غير موجود في ورقة قاعدة بيانات الأسر")

    data = []
    fam_tot = dict(core=0, ext=0, total=0, casesExt=0, sub=0, dis=0, eld=0, chr=0, war=0, nurse=0)
    for code in codes:
        completion = num(summ.get(code, c_comp))
        p1 = num(summ.get(code, c_p1))
        p2 = num(summ.get(code, c_p2))
        sched = norm(summ.get(code, c_sched))
        elig = norm(summ.get(code, c_elig)) == "نعم"
        has_ext = norm(fam.get(code, a_hasext)) == "نعم"
        rec = dict(
            code=code,
            name=norm(summ.get(code, c_name)),
            eng=ENG_MAP[code],
            items=int(num(summ.get(code, c_items))),
            completion=round(completion, 10), p1=round(p1, 10), p2=round(p2, 10),
            schedGood=sched.startswith("🟢"),
            eligGood=elig,
            housing=norm(summ.get(code, c_housing)),
            schedText=sched,
            eligText=("✅ مؤهل لصرف الدفعة الثالثة - تم إنجاز كافة البنود" if elig else
                      f"⛔ غير مؤهل - الدفعة الثالثة تُصرف بعد إنجاز 100% من جدول الكميات (المتبقي {round((1-completion)*100)}%)"),
            stage=norm(flow.get(code, f_stage)),
            adherence=norm(flow.get(code, f_adh)),
            familySize=int(num(fam.get(code, a_total))),
            hasExtended=has_ext,
            disabled=int(num(fam.get(code, a_dis))),
            chronic=int(num(fam.get(code, a_chr))),
            warInjured=int(num(fam.get(code, a_war))),
            elderly=int(num(fam.get(code, a_eld))),
        )
        data.append(rec)
        fam_tot["core"] += num(fam.get(code, a_core))
        fam_tot["ext"] += num(fam.get(code, a_extm))
        fam_tot["total"] += rec["familySize"]
        fam_tot["casesExt"] += 1 if has_ext else 0
        fam_tot["sub"] += num(fam.get(code, a_subfam))
        fam_tot["dis"] += rec["disabled"]
        fam_tot["eld"] += rec["elderly"]
        fam_tot["chr"] += rec["chronic"]
        fam_tot["war"] += rec["warInjured"]
        fam_tot["nurse"] += num(fam.get(code, a_nurse))

    # ---- الفئات العمرية
    ws_age = wb["الفئات العمرية للمشروع"]
    pos = find_label(ws_age, "الفئة العمرية")
    if not pos:
        raise ValidationError("لم يُعثر على جدول الفئات العمرية")
    r0, c0 = pos
    age_bands = []
    for r in range(r0 + 1, r0 + 5):
        label = norm(ws_age.cell(r, c0).value)
        male, female, total = (int(num(ws_age.cell(r, c0 + k).value)) for k in (1, 2, 3))
        pct = num(ws_age.cell(r, c0 + 4).value)
        age_bands.append(dict(label=label, male=male, female=female, total=total,
                              pct=round(pct * 100 if pct <= 1 else pct, 1)))
    if [b["label"] for b in age_bands] != ["أقل من 8 سنوات", "8 - 17 سنة", "18 - 59 سنة", "60 سنة فأكثر"]:
        raise ValidationError("ترتيب أو أسماء الفئات العمرية في المصدر تغيّر: " + str([b["label"] for b in age_bands]))

    # ---- متوسط أفراد الأسرة من لوحة المؤشرات
    ws_kpi = wb["لوحة المؤشرات"]
    p = find_label(ws_kpi, "متوسط أفراد الأسرة الواحدة")
    if not p:
        raise ValidationError("لم يُعثر على خانة متوسط أفراد الأسرة الواحدة")
    avg_family = round(num(ws_kpi.cell(p[0] + 1, p[1]).value), 1)

    individuals = int(fam_tot["total"])
    children = age_bands[0]["total"] + age_bands[1]["total"]
    women = sum(b["female"] for b in age_bands)
    men = sum(b["male"] for b in age_bands)
    family = dict(
        families=len(data), coreMembers=int(fam_tot["core"]), extendedMembers=int(fam_tot["ext"]),
        totalIndividuals=individuals, casesWithExtended=int(fam_tot["casesExt"]),
        extendedSubfamilies=int(fam_tot["sub"]), disabled=int(fam_tot["dis"]), elderly=int(fam_tot["eld"]),
        chronic=int(fam_tot["chr"]), warInjured=int(fam_tot["war"]), nursingMothers=int(fam_tot["nurse"]),
        avgFamilySize=avg_family,
        childrenUnder18=children, childrenPct=round(children / individuals * 100, 1) if individuals else 0,
        women=women, men=men, workingAge=age_bands[2]["total"], workingAgePct=age_bands[2]["pct"],
        elderlyPct=age_bands[3]["pct"],
    )

    # ---- المالية (مستوى المشروع فقط)
    paid = [round(sum(num(summ.get(c, col)) for c in codes), 2) for col in c_paid]
    paid_total_rows = round(sum(num(summ.get(c, c_paid_total)) for c in codes), 2)
    contracts_sum = round(sum(num(flow.get(c, f_contract)) for c in codes), 2)

    ws_flow = flow.ws
    pt = find_label(ws_flow, "الدفعة", rows=range(max(flow.rows.values()) + 1, ws_flow.max_row + 1))
    planned, earned = [], []
    labels = []
    if pt:
        for r in range(pt[0] + 1, pt[0] + 4):
            labels.append(norm(ws_flow.cell(r, pt[1]).value))
            planned.append(round(num(ws_flow.cell(r, pt[1] + 1).value), 2))
            earned.append(round(num(ws_flow.cell(r, pt[1] + 2).value), 2))
    if len(planned) != 3 or not all(l.startswith("الدفعة") for l in labels):
        raise ValidationError("لم يُعثر على جدول (الدفعة | القيمة المخطط لها | المستحق فعليًا) أسفل لوحة سير المشروع")

    receipts = []
    for col in c_rcv:
        dates = sorted({d for d in (as_date(summ.get(c, col)) for c in codes) if d})
        got = [c for c in codes if as_date(summ.get(c, col))]
        receipts.append(dict(received=len(got), dates=dates, missing=[c for c in codes if c not in got]))

    total_disbursed = round(sum(paid), 2)
    finance = dict(
        totalGrant=TOTAL_GRANT,
        totalDisbursed=total_disbursed,
        remaining=round(TOTAL_GRANT - total_disbursed, 2),
        disbursedPct=round(total_disbursed / TOTAL_GRANT * 100, 1),
        contractsSum=contracts_sum,
        payments=[dict(label=labels[i], planned=planned[i], actual=paid[i]) for i in range(3)],
        earned=dict(total=round(sum(earned), 2), payments=earned),
        receipts=receipts,
    )
    # قيم الهوية والجوال الفعلية من المصدر - فقط للتأكد أنها لم تتسرّب إلى الناتج
    sensitive = set()
    for col in (summ.col("رقم الهوية"), summ.col("رقم الجوال")):
        for c in codes:
            v = summ.get(c, col)
            if v not in (None, ""):
                t = str(int(v)) if isinstance(v, float) and v.is_integer() else norm(v)
                if len(t) >= 7:
                    sensitive.add(t)
    checks = dict(paid_total_rows=paid_total_rows, sensitive=sensitive)
    return data, family, finance, age_bands, checks


# ---------------------------------------------------------------- كتل JS داخل HTML
DATA_START = "/* ============================= DATA"
DATA_END_RE = re.compile(r"^/\* =+ \*/\s*$", re.M)


def data_region(html):
    s = html.find(DATA_START)
    if s < 0:
        raise ValidationError("لم يُعثر على بداية كتلة البيانات في HTML")
    m = DATA_END_RE.search(html, s + len(DATA_START))
    if not m:
        raise ValidationError("لم يُعثر على نهاية كتلة البيانات في HTML")
    return s, m.end()


def block_span(html, name, start, end):
    m = re.compile(r"^const " + re.escape(name) + r" = ", re.M).search(html, start, end)
    if not m:
        return None
    line_end = html.find("\n", m.start())
    first_line = html[m.start():line_end].rstrip()
    if first_line.endswith(";"):
        return m.start(), line_end
    close = re.compile(r"^[\]\}];[ \t]*$", re.M).search(html, line_end, end)
    if not close:
        raise ValidationError(f"كتلة {name} غير مغلقة")
    return m.start(), close.end()


def replace_block(html, name, new_text):
    s, e = data_region(html)
    span = block_span(html, name, s, e)
    if span is None:
        anchor = block_span(html, "DATA", s, e)
        return html[:anchor[0]] + new_text + "\n\n" + html[anchor[0]:]
    return html[:span[0]] + new_text + html[span[1]:]


def read_old_state(html):
    """يقيّم كتلة البيانات الحالية بـ node ليقرأ القيم القديمة بدقة (لا تحليل نصي يدوي)."""
    s, e = data_region(html)
    js = html[s:e] + """
;console.log(JSON.stringify({
  REPORT_DATE, DATA, HISTORY,
  PREV: (typeof PREV!=="undefined") ? PREV : null,
  SYNC: (typeof SYNC!=="undefined") ? SYNC : null,
  MILESTONES, COORDS
}));"""
    with tempfile.NamedTemporaryFile("w", suffix=".mjs", delete=False, encoding="utf-8") as f:
        f.write(js)
        path = f.name
    try:
        out = subprocess.run(["node", path], capture_output=True, text=True, check=True).stdout
    finally:
        os.unlink(path)
    return json.loads(out)


def js(v):
    return json.dumps(v, ensure_ascii=False, separators=(",", ":"))


def fmt_data(data):
    return "const DATA = [\n" + "".join(js(d) + ",\n" for d in data) + "];"


def score(d):
    return round(50 * d["completion"] + 20 * ((d["p1"] + d["p2"]) / 2) + (30 if d["schedGood"] else 0))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--html", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--file-id", default="")
    ap.add_argument("--now", help="وقت المزامنة بتوقيت غزة YYYY-MM-DDTHH:MM (للاختبار)")
    a = ap.parse_args()

    if a.now:
        now_gaza = datetime.strptime(a.now, "%Y-%m-%dT%H:%M")
    else:
        try:
            from zoneinfo import ZoneInfo
            now_gaza = datetime.now(ZoneInfo("Asia/Gaza")).replace(tzinfo=None)
        except Exception:
            now_gaza = (datetime.now(timezone.utc) + GAZA_OFFSET).replace(tzinfo=None)
    today = now_gaza.date().isoformat()

    html = open(a.html, encoding="utf-8").read()
    old = read_old_state(html)
    data, family, finance, age_bands, checks = read_source(a.xlsx)

    # ---- التحقق
    errs = []
    if len(data) != EXPECTED_N:
        errs.append(f"عدد المستفيدين {len(data)} وليس {EXPECTED_N}")
    for d in data:
        if set(d) != set(FIELDS):
            errs.append(f"حقول ناقصة/زائدة في {d['code']}")
        for k in ("completion", "p1", "p2"):
            if not (0 <= d[k] <= 1.0000001):
                errs.append(f"{d['code']}.{k} خارج النطاق 0..1: {d[k]}")
        if not d["name"] or not d["stage"]:
            errs.append(f"{d['code']}: اسم أو مرحلة فارغة")
    if family["totalIndividuals"] != family["women"] + family["men"]:
        errs.append(f"الإناث+الذكور ({family['women']}+{family['men']}) لا يساوي إجمالي الأفراد {family['totalIndividuals']}")
    if sum(b["total"] for b in age_bands) != family["totalIndividuals"]:
        errs.append("مجموع الفئات العمرية لا يساوي إجمالي الأفراد")
    if abs(checks["paid_total_rows"] - finance["totalDisbursed"]) > 0.05:
        errs.append(f"عمود إجمالي المصروف ({checks['paid_total_rows']}) لا يطابق مجموع الدفعات الثلاث ({finance['totalDisbursed']})")
    if any(p["actual"] < 0 or p["planned"] <= 0 for p in finance["payments"]):
        errs.append("قيم دفعات سالبة أو مخطط صفري")
    if finance["totalDisbursed"] > TOTAL_GRANT:
        errs.append("المصروف أكبر من المنحة")
    if errs:
        raise ValidationError(" | ".join(errs))

    # ---- السجل التاريخي: نقطة واحدة لكل يوم (آخر مزامنة في اليوم تغلب)
    avg_score = round(sum(score(d) for d in data) / len(data), 1)
    avg_comp = round(sum(d["completion"] for d in data) / len(data) * 100, 1)
    history = [h for h in old["HISTORY"] if h["date"] != today]
    history.append(dict(date=today, overall=avg_score, completion=avg_comp))
    history = sorted(history, key=lambda h: h["date"])[-120:]

    # ---- اللقطة السابقة للمقارنة: آخر لقطة بتاريخ يوم سابق
    prev = old["PREV"]
    if old["REPORT_DATE"] != today:
        prev = dict(date=old["REPORT_DATE"],
                    rows={d["code"]: [round(d["completion"], 4), round(d["p1"], 4), round(d["p2"], 4), d["stage"]]
                          for d in old["DATA"]})

    old_by = {d["code"]: d for d in old["DATA"]}
    changed = [d["code"] for d in data if d["code"] not in old_by or any(
        abs(d[k] - old_by[d["code"]][k]) > 1e-6 for k in ("completion", "p1", "p2")) or d["stage"] != old_by[d["code"]]["stage"]]

    sync = dict(at=now_gaza.strftime("%Y-%m-%d %H:%M"), tz="توقيت فلسطين", fileId=a.file_id or None)

    out = html
    out = replace_block(out, "REPORT_DATE", f'const REPORT_DATE = "{today}";')
    out = replace_block(out, "SYNC", "const SYNC = " + js(sync) + ";")
    out = replace_block(out, "HISTORY", "const HISTORY = [\n" + "".join("  " + js(h) + ",\n" for h in history) + "];")
    out = replace_block(out, "FAMILY", "const FAMILY = " + js(family) + ";")
    out = replace_block(out, "FINANCE", "const FINANCE = " + json.dumps(finance, ensure_ascii=False, indent=2) + ";")
    out = replace_block(out, "AGE_BANDS", "const AGE_BANDS = [\n" + "".join("  " + js(b) + ",\n" for b in age_bands) + "];")
    out = replace_block(out, "PREV", "const PREV = " + js(prev) + ";")
    out = replace_block(out, "DATA", fmt_data(data))

    # ---- تحقق ما بعد الاستبدال
    new_state = read_old_state(out)
    if new_state["MILESTONES"] != old["MILESTONES"] or new_state["COORDS"] != old["COORDS"]:
        raise ValidationError("تغيّرت كتلة MILESTONES أو COORDS — ممنوع في التحديث التلقائي")
    if len(new_state["DATA"]) != EXPECTED_N or new_state["REPORT_DATE"] != today:
        raise ValidationError("الناتج لا يحوي البيانات الجديدة كما يجب")
    if not checks["sensitive"]:
        raise ValidationError("تعذّر قراءة عمودي الهوية/الجوال للتحقق من عدم تسرّبهما")
    if any(t in out for t in checks["sensitive"]):
        raise ValidationError("رقم هوية أو جوال من الملف المصدر ظهر في الناتج")
    if "Ã" in out or "Ø" in out:
        raise ValidationError("ترميز تالف (mojibake) في الناتج")
    scripts = re.findall(r"<script>([\s\S]*?)</script>", out)
    with tempfile.NamedTemporaryFile("w", suffix=".js", delete=False, encoding="utf-8") as f:
        f.write("\n;\n".join(scripts))
        path = f.name
    try:
        r = subprocess.run(["node", "--check", path], capture_output=True, text=True)
        if r.returncode != 0:
            raise ValidationError("خطأ نحوي في JavaScript الناتج: " + r.stderr[:500])
    finally:
        os.unlink(path)

    open(a.out, "w", encoding="utf-8").write(out)
    result = dict(
        ok=True, reportDate=today, syncAt=sync["at"], fileId=a.file_id or None,
        avgScore=avg_score, avgCompletion=avg_comp,
        disbursed=finance["totalDisbursed"], disbursedPct=finance["disbursedPct"],
        prevDate=(prev or {}).get("date"), changedSinceLastPublish=changed,
    )
    print("RESULT:" + json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except ValidationError as ex:
        print("RESULT:" + json.dumps(dict(ok=False, error=str(ex)), ensure_ascii=False))
        sys.exit(2)
